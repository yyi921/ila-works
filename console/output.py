"""What a command returns: a titled list of blocks (headline figures, tables, notes), rendered as JSON for the
console or written to CSV/XLSX for ``>>``."""

import csv
import io
from dataclasses import dataclass, field
from datetime import date
from decimal import Decimal

MONEY, PCT, TEXT, DATE, NUMBER = "money", "pct", "text", "date", "number"


def fmt_money(value):
    if value is None:
        return ""
    value = Decimal(value).quantize(Decimal("0.01"))
    text = f"{abs(value):,.2f}"
    return f"({text})" if value < 0 else text


def fmt_pct(value):
    if value is None:
        return ""
    value = Decimal(value) * 100
    text = f"{abs(value):,.1f}%"
    return f"({text})" if value < 0 else text


def fmt_date(value):
    return value.strftime("%d %b %Y").lstrip("0") if isinstance(value, date) else (value or "")


def fmt(value, kind):
    if kind == MONEY:
        return fmt_money(value)
    if kind == PCT:
        return fmt_pct(value)
    if kind == DATE:
        return fmt_date(value)
    if kind == NUMBER:
        return "" if value is None else f"{value:,}"
    return "" if value is None else str(value)


@dataclass
class Column:
    label: str
    kind: str = TEXT


@dataclass
class Table:
    columns: list
    rows: list = field(default_factory=list)
    title: str = ""
    # Indexes of rows to show as totals.
    total_rows: set = field(default_factory=set)

    def add(self, *cells, total=False):
        if total:
            self.total_rows.add(len(self.rows))
        self.rows.append(list(cells))

    def to_json(self):
        return {
            "type": "table", "title": self.title,
            "columns": [{"label": c.label, "kind": c.kind} for c in self.columns],
            "rows": [{"cells": [fmt(v, c.kind) for v, c in zip(row, self.columns)],
                      "negative": [isinstance(v, Decimal) and v < 0 for v in row],
                      "total": i in self.total_rows} for i, row in enumerate(self.rows)],
        }


@dataclass
class Figure:
    label: str
    value: object
    kind: str = MONEY
    detail: str = ""


@dataclass
class Figures:
    items: list

    def to_json(self):
        return {"type": "figures", "items": [
            {"label": f.label, "value": fmt(f.value, f.kind), "detail": f.detail,
             "negative": isinstance(f.value, Decimal) and f.value < 0} for f in self.items]}


@dataclass
class Note:
    text: str
    level: str = "info"  # info, warning

    def to_json(self):
        return {"type": "note", "text": self.text, "level": self.level}


@dataclass
class Result:
    title: str
    blocks: list = field(default_factory=list)
    subtitle: str = ""
    link: str = ""

    def add(self, block):
        self.blocks.append(block)
        return block

    def to_json(self):
        return {"title": self.title, "subtitle": self.subtitle, "link": self.link,
                "blocks": [b.to_json() for b in self.blocks]}

    def tables(self):
        return [b for b in self.blocks if isinstance(b, Table)]


class ExportError(ValueError):
    pass


def _figures_table(result):
    figures = [f for b in result.blocks if isinstance(b, Figures) for f in b.items]
    if not figures:
        return None
    table = Table([Column("Measure"), Column("Value"), Column("Detail")], title="Summary")
    for f in figures:
        table.add(f.label, f.value, f.detail)
    return table


def export(result, filename):
    """Bytes and content type for ``>> filename``; the extension picks the format (.csv or .xlsx)."""
    tables = result.tables()
    summary = _figures_table(result)
    if summary:
        tables = [summary] + tables
    if not tables:
        raise ExportError(f"{result.title} has no table to save to a file.")
    lower = filename.lower()
    if lower.endswith(".csv"):
        buffer = io.StringIO()
        writer = csv.writer(buffer)
        for i, table in enumerate(tables):
            if i:
                writer.writerow([])
            if len(tables) > 1 and table.title:
                writer.writerow([table.title])
            writer.writerow([c.label for c in table.columns])
            for row in table.rows:
                writer.writerow(["" if v is None else v.isoformat() if isinstance(v, date) else v for v in row])
        return buffer.getvalue().encode("utf-8-sig"), "text/csv"
    if lower.endswith(".xlsx"):
        return _xlsx(result, tables), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    raise ExportError("Save as .xlsx or .csv, e.g. >> report.xlsx.")


def _xlsx(result, tables):
    from openpyxl import Workbook
    from openpyxl.styles import Font

    book = Workbook()
    book.remove(book.active)
    used = set()
    for i, table in enumerate(tables, 1):
        name = (table.title or f"Sheet {i}")[:31]
        for ch in "[]:*?/\\":
            name = name.replace(ch, " ")
        while name in used:
            name = f"{name[:28]} {i}"
        used.add(name)
        sheet = book.create_sheet(name)
        sheet.append([result.title])
        sheet["A1"].font = Font(bold=True)
        if result.subtitle:
            sheet.append([result.subtitle])
        sheet.append([])
        sheet.append([c.label for c in table.columns])
        header_row = sheet.max_row
        for cell in sheet[header_row]:
            cell.font = Font(bold=True)
        for r, row in enumerate(table.rows):
            values = []
            for v, col in zip(row, table.columns):
                if col.kind in (MONEY, PCT, NUMBER) and v is not None:
                    values.append(float(v))
                elif col.kind == TEXT and v is not None:
                    values.append(str(v))  # text-typed so codes like 0400 keep their zeros
                else:
                    values.append(v)
            sheet.append(values)
            for cell, col in zip(sheet[sheet.max_row], table.columns):
                if col.kind == MONEY:
                    cell.number_format = "#,##0.00;(#,##0.00)"
                elif col.kind == PCT:
                    cell.number_format = "0.0%;(0.0%)"
                elif col.kind == DATE:
                    cell.number_format = "d mmm yyyy"
                elif col.kind == TEXT:
                    cell.number_format = "@"
                if r in table.total_rows:
                    cell.font = Font(bold=True)
        for idx, col in enumerate(table.columns, 1):
            width = max([len(col.label)] + [len(fmt(row[idx - 1], col.kind)) for row in table.rows]) + 2
            sheet.column_dimensions[sheet.cell(row=header_row, column=idx).column_letter].width = min(width, 60)
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()
