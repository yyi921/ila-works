import io

import pytest
from django.contrib.auth import get_user_model
from django.contrib.auth.models import Permission
from openpyxl import load_workbook

from console import builder
from console.models import CustomCommand
from console.parser import parse
from core.models import AuditLog, Organisation

from .fake_xero import fake_source


@pytest.fixture
def staff(db):
    return get_user_model().objects.create_user("builder", password="pw")


@pytest.fixture
def orgs(staff):
    demo = Organisation.objects.create(name="Demo Events")
    other = Organisation.objects.create(name="Other Venue")
    demo.members.add(staff)
    other.members.add(staff)
    return demo, other


@pytest.fixture
def can_build(staff):
    staff.user_permissions.add(Permission.objects.get(codename="build_commands"))
    return staff


@pytest.fixture
def client(client, staff):
    client.force_login(staff)
    return client


@pytest.fixture(autouse=True)
def fake(monkeypatch):
    source = fake_source()
    monkeypatch.setattr("console.views.source_for", lambda organisation: source)
    return source


def make(org, name="_month", steps=("/status --September 2026 --revenue", "/bills"), **extra):
    return CustomCommand.objects.create(organisation=org, name=name, summary="Month pack",
                                        definition={"steps": list(steps)}, **extra)


def run(client, command):
    return client.post("/console/run/", {"command": command}).json()


# --- Syntax and validation --------------------------------------------------------------------------------------

def test_built_commands_parse_with_or_without_slash():
    assert parse("/_month --mtd").name == "_month"
    cmd = parse("_supplier onemusic --Oct 2024 to Sep 2026")
    assert (cmd.name, cmd.text, cmd.params) == ("_supplier", "onemusic", ["Oct 2024 to Sep 2026"])


@pytest.mark.parametrize("raw, expected", [("weekly", "_weekly"), ("/_Weekly", "_Weekly"), ("_m-2", "_m-2"), ("createInvoice", "_createInvoice")])
def test_clean_name_adds_underscore(raw, expected):
    assert builder.clean_name(raw) == expected


@pytest.mark.parametrize("raw", ["_", "_1abc", "_has space", "_" + "a" * 40])
def test_clean_name_rejects(raw):
    with pytest.raises(ValueError):
        builder.clean_name(raw)


def test_clean_steps_accepts_native_commands_and_placeholders():
    assert builder.clean_steps("\n/compare --bill {text} {params}\n\n/creditors\n") == [
        "/compare --bill {text} {params}", "/creditors"]


@pytest.mark.parametrize("steps, problem", [
    ("", "at least one"),
    ("/nonsense", "no /nonsense command"),
    ("/use Demo", "/use can't be used"),
    ("/_other", "can't run other built commands"),
    ("/recon << statement.csv", "leave out << and >>"),
    ("/bills >> bills.xlsx", "leave out << and >>"),
    ("/inv {number}", "{number} isn't a placeholder"),
    ("status", 'start with "/"'),
    ("\n".join(["/bills"] * 11), "at most 10"),
])
def test_clean_steps_rejects(steps, problem):
    with pytest.raises(ValueError) as caught:
        builder.clean_steps(steps)
    assert any(problem in message for message in caught.value.args)


def test_clean_steps_lists_every_problem():
    with pytest.raises(ValueError) as caught:
        builder.clean_steps("/status\n/use Demo\n/inv {number}")
    assert len(caught.value.args) == 2 and caught.value.args[0].startswith("Line 2")


# --- Builder page -----------------------------------------------------------------------------------------------

def test_builder_needs_permission(client, orgs):
    assert client.get("/console/builder/").status_code == 403
    assert client.post("/console/builder/new/shortcut/", {}).status_code == 403
    assert b"Builder" not in client.get("/console/").content.split(b"<main")[0]


def test_builder_lists_templates_and_links_from_console(client, orgs, can_build):
    page = client.get("/console/builder/").content.decode()
    assert "Start from scratch" in page and "_createInvoice" in page and "Not built yet" not in page
    assert 'href="/console/builder/"' in client.get("/console/").content.decode()


def test_superuser_can_build_without_the_permission(client, orgs, staff):
    staff.is_superuser = True
    staff.save()
    assert client.get("/console/builder/").status_code == 200


def test_create_invoice_template_makes_a_command_without_steps(client, orgs, can_build):
    page = client.get("/console/builder/new/createinvoice/").content.decode()
    assert 'value="_createInvoice"' in page and "Invoice set-ups" in page
    client.post("/console/builder/new/createinvoice/", {"name": "_createInvoice", "summary": "Invoices",
                                                        "is_active": "on"})
    custom = CustomCommand.objects.get()
    assert (custom.name, custom.template, custom.definition) == ("_createInvoice", "createinvoice", {})
    page = client.post("/console/builder/new/shortcut/", {"name": "_CREATEINVOICE", "summary": "x",
                                                         "steps": "/bills"}).content.decode()
    assert "already has a command called _CREATEINVOICE" in page
    assert client.get("/console/builder/new/nothing/").status_code == 404


def test_create_edit_delete_are_audited(client, orgs, can_build):
    resp = client.post("/console/builder/new/shortcut/", {
        "name": "month", "summary": "Month pack", "steps": "/status --mtd\n/bills", "is_active": "on"})
    assert resp.status_code == 302
    custom = CustomCommand.objects.get()
    assert (custom.name, custom.steps, custom.organisation, custom.created_by) == (
        "_month", ["/status --mtd", "/bills"], orgs[0], can_build)

    client.post(f"/console/builder/{custom.pk}/", {"name": "_month", "summary": "Month pack v2", "steps": "/bills"})
    custom.refresh_from_db()
    assert custom.summary == "Month pack v2" and custom.steps == ["/bills"] and not custom.is_active

    assert client.get(f"/console/builder/{custom.pk}/delete/").status_code == 200
    client.post(f"/console/builder/{custom.pk}/delete/")
    assert not CustomCommand.objects.exists()

    logs = list(AuditLog.objects.filter(action=AuditLog.Action.BUILD).order_by("created_at"))
    assert [log.input for log in logs] == ["create _month", "update _month", "delete _month"]
    assert logs[1].result["before"]["summary"] == "Month pack" and logs[1].result["after"]["summary"] == "Month pack v2"
    assert logs[2].result["after"] is None


def test_form_errors_are_shown(client, orgs, can_build):
    make(orgs[0])
    resp = client.post("/console/builder/new/shortcut/", {"name": "_month", "summary": "x", "steps": "/use Other"})
    page = resp.content.decode()
    assert resp.status_code == 200 and "already has a command called _month" in page and "/use can" in page


def test_other_organisations_commands_are_out_of_reach(client, orgs, can_build):
    theirs = make(orgs[1])
    assert client.get(f"/console/builder/{theirs.pk}/").status_code == 404
    assert client.post(f"/console/builder/{theirs.pk}/delete/").status_code == 404
    assert "_month" not in client.get("/console/builder/").content.decode()


# --- Running built commands -------------------------------------------------------------------------------------

def test_shortcut_runs_its_steps_as_one_report(client, orgs):
    make(orgs[0])
    data = run(client, "_month")
    assert data["ok"], data
    result = data["result"]
    assert result["title"] == "/_month" and result["subtitle"] == "Month pack"
    sections = [b for b in result["blocks"] if b["type"] == "section"]
    assert [s["command"] for s in sections] == ["/status --September 2026 --revenue", "/bills"]
    assert AuditLog.objects.get(action=AuditLog.Action.COMMAND).input == "_month"


def test_placeholders_take_text_and_params(client, orgs, monkeypatch):
    make(orgs[0], "_pl", ["/pl {params}"])
    make(orgs[0], "_doc", ["/inv {text}"])
    data = run(client, "/_pl --September 2026")
    assert [b["command"] for b in data["result"]["blocks"] if b["type"] == "section"] == ["/pl --September 2026"]
    assert "needs text" in run(client, "/_doc")["error"]
    assert "doesn't take text" in run(client, "/_pl oops")["error"]
    assert "doesn't take -- parameters" in run(client, "/_doc INV-1 --mtd")["error"]


def test_a_failing_step_is_reported_and_the_rest_still_run(client, orgs):
    make(orgs[0], "_mixed", ["/inv NOPE-404", "/bills"])
    blocks = run(client, "/_mixed")["result"]["blocks"]
    assert blocks[0] == {"type": "section", "command": "/inv NOPE-404", "title": "Couldn't run", "subtitle": ""}
    assert blocks[1]["type"] == "note" and blocks[1]["level"] == "warning"
    assert any(b["type"] == "section" and b["command"] == "/bills" for b in blocks)


def test_inactive_and_other_organisations_commands_do_not_run(client, orgs):
    make(orgs[0], "_off", is_active=False)
    make(orgs[1], "_theirs")
    assert "no /_off command for Demo Events" in run(client, "/_off")["error"]
    assert "no /_theirs command" in run(client, "/_theirs")["error"]


def test_built_command_saves_every_table(client, orgs):
    make(orgs[0])
    resp = client.post("/console/run/", {"command": "/_month >> pack.xlsx"})
    book = load_workbook(io.BytesIO(resp.content))
    assert len(book.sheetnames) >= 2


def test_help_lists_and_explains_built_commands(client, orgs):
    make(orgs[0], "_doc", ["/inv {text}"])
    blocks = run(client, "/help")["result"]["blocks"]
    assert blocks[-1]["title"] == "Demo Events's own commands"
    assert blocks[-1]["rows"][0]["cells"] == ["/_doc", "Month pack"]
    detail = run(client, "/help _doc")["result"]
    assert detail["title"] == "/_doc" and detail["blocks"][1]["rows"][0]["cells"] == ["1", "/inv {text}"]


def test_console_completes_and_prefills_built_commands(client, orgs):
    make(orgs[0])
    page = client.get("/console/?cmd=/_month").content.decode()
    assert '"_month"' in page and 'value="/_month"' in page
