import csv
import io
from decimal import Decimal

import pytest
from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile

from console.models import SplitRule, SplitRun
from core.models import Assignment, Dimension, DimensionValue, JournalLine, Organisation
from core.tagging import unbalanced_lines
from xero.source import ProfitAndLoss

from .fake_xero import fake_source

D = Decimal
JOURNALS = [
    {"ManualJournalID": "mj-wages", "Status": "POSTED", "DateString": "2026-09-30T00:00:00", "Narration": "Payroll",
     "JournalLines": [{"Description": "Wages", "AccountCode": "477", "LineAmount": "400.00"},
                      {"Description": "Rent share", "AccountCode": "469", "LineAmount": "200.00"},
                      {"Description": "Accrual", "AccountCode": "800", "LineAmount": "-600.00"}]},
]
INVOICES = [
    {"InvoiceID": "b-cons", "InvoiceNumber": "B-CONS", "Type": "ACCPAY", "Status": "AUTHORISED",
     "DateString": "2026-09-15T00:00:00", "Contact": {"ContactID": "c-adv", "Name": "Advisers"},
     "LineAmountTypes": "Exclusive", "LineItems": [
         {"LineItemID": "bc-1", "Description": "Consulting", "AccountCode": "412", "LineAmount": "1000.00"}]},
    {"InvoiceID": "b-wage", "InvoiceNumber": "B-WAGE", "Type": "ACCPAY", "Status": "AUTHORISED",
     "DateString": "2026-09-16T00:00:00", "Contact": {"ContactID": "c-hire", "Name": "Labour Hire"},
     "LineAmountTypes": "Exclusive", "LineItems": [
         {"LineItemID": "bw-1", "Description": "Casual staff", "AccountCode": "477", "LineAmount": "250.00"}]},
]


@pytest.fixture
def org(db):
    user = get_user_model().objects.create_user("finance", password="pw")
    org = Organisation.objects.create(name="TECC")
    org.members.add(user)
    return org, user


@pytest.fixture
def client(client, org):
    client.force_login(org[1])
    return client


@pytest.fixture(autouse=True)
def source(monkeypatch):
    source = fake_source()
    source.raw_invoices = lambda start, end: INVOICES
    source.raw_credit_notes = lambda start, end: []
    source.raw_bank_transactions = lambda start, end: []
    source.raw_manual_journals = lambda start, end: JOURNALS
    source.profit_and_loss = lambda start, end: ProfitAndLoss(
        by_account={"a-wages": D("650"), "a-rent": D("200"), "a-cons": D("1000")},
        category_of={"a-wages": "expenses", "a-rent": "expenses", "a-cons": "expenses"}, net_profit=D("-1850"))
    monkeypatch.setattr("console.views.source_for", lambda organisation: source)
    return source


@pytest.fixture
def event(org):
    dim = Dimension.objects.create(organisation=org[0], code="event", name="Event")
    for code, name in [("ABBA", "ABBA"), ("ADF", "ADF Seminar"), ("CN", "Celtic Nights")]:
        DimensionValue.objects.create(organisation=org[0], dimension=dim, code=code, name=name)
    return dim


def say(client, text, files=()):
    data = client.post("/console/run/", {"command": text, "file": list(files)}).json()
    assert data["ok"], data
    return data["result"]


def error(client, text, files=()):
    data = client.post("/console/run/", {"command": text, "file": list(files)}).json()
    assert not data["ok"], data
    return data["error"]


def notes(result):
    return " ".join(b["text"] for b in result["blocks"] if b["type"] == "note")


def tables(result):
    return {b["title"]: [r["cells"] for r in b["rows"]] for b in result["blocks"] if b["type"] == "table"}


HEAD = ["Rule", "Accounts", "Types", "Contact", "Description", "From", "To", "Dimension", "Value", "Amount", "Percent"]


def rules_file(rows, name="rules.csv"):
    out = io.StringIO()
    csv.writer(out).writerows([HEAD] + rows)
    return SimpleUploadedFile(name, out.getvalue().encode("utf-8-sig"), content_type="text/csv")


LABOUR = [["labour", "469-477", "journals", "", "", "", "", "Event", "ABBA", "100", ""],
          ["labour", "469-477", "journals", "", "", "", "", "Event", "ADF", "300", ""]]
HYBRID = [["advice", "412", "bills", "", "", "", "", "Event", "ABBA", "100", ""],
          ["advice", "412", "bills", "", "", "", "", "Event", "ADF", "", "30"],
          ["advice", "412", "bills", "", "", "", "", "Event", "CN", "", "70"]]


def test_amounts_are_pro_rated_across_a_group_of_accounts(client, org, event):
    result = say(client, "/splittag --Sep 2026 <<", [rules_file(LABOUR + HYBRID)])
    t = tables(result)
    labour = t["labour: 469-477 · journals · 2 line(s) · September 2026"]
    assert labour == [["ABBA", "100.00", "16.7%", "100.00"], ["ADF ADF Seminar", "300.00", "50.0%", "300.00"],
                      ["Unassigned", "What's left", "33.3%", "200.00"], ["Total", "", "100.0%", "600.00"]]
    advice = t["advice: 412 · bills · 1 line(s) · September 2026"]
    assert advice[:3] == [["ABBA", "100.00", "10.0%", "100.00"],
                          ["ADF ADF Seminar", "30% of what's left", "27.0%", "270.00"],
                          ["CN Celtic Nights", "70% of what's left", "63.0%", "630.00"]]
    assert advice[3][3] == "0.00"
    assert not Assignment.objects.exists()  # nothing before approval

    result = say(client, "approve")
    assert result["title"] == "Split saved: 3 line(s) tagged"
    wages = JournalLine.objects.get(description="Wages")
    assert sorted((t.dimension_value.code, t.allocation) for t in wages.tags.all()) == [
        ("ABBA", D("0.166675")), ("ADF", D("0.5"))]  # $66.67 and $200.00 of the $400 line, to the cent
    assert unbalanced_lines(org[0], event) == []
    assert {r.name: r.last_run for r in SplitRule.objects.all()} == {"labour": "September 2026",
                                                                     "advice": "September 2026"}

    pl = next(b for b in say(client, "/pl --event --Sep 2026")["blocks"] if b["type"] == "table")
    rows = {r["cells"][1]: r["cells"][2:] for r in pl["rows"]}
    assert [c["label"] for c in pl["columns"]][2:6] == ["ABBA", "ADF ADF Seminar", "CN Celtic Nights", "Unassigned"]
    assert rows["Wages"][:4] == ["66.67", "200.00", "0.00", "383.33"]  # the bill's 250 isn't in the journals rule
    assert rows["Rent"][:4] == ["33.33", "100.00", "0.00", "66.67"]
    assert rows["Consulting"][:4] == ["100.00", "270.00", "630.00", "0.00"]


def test_saved_rule_runs_again_with_the_scope_changed(client, org, event):
    say(client, "/splittag --Sep 2026 <<", [rules_file(LABOUR)])
    say(client, "approve")
    listing = tables(say(client, "/splittag"))["Saved rules"]
    assert listing == [["labour", "Event", "469-477", "journals", "ABBA $100, ADF $300", "September 2026"]]

    result = say(client, "/splittag --gl 477 --sep-2026 --labour --all")  # wages journal line + the casual bill
    table = tables(result)["labour: 477 · all types · 2 line(s) · September 2026"]
    assert table[-1] == ["Total", "", "100.0%", "650.00"]
    assert "already have a split" in notes(result)
    say(client, "approve")
    rule = SplitRule.objects.get()
    assert rule.scope["accounts"] == [["469", "477"]]  # the saved scope isn't changed by --gl
    wages = JournalLine.objects.get(description="Wages")
    assert sorted(t.allocation for t in wages.tags.exclude(status__in=("superseded", "rejected"))) == [
        D("0.15385"), D("0.46155")]  # $61.54 and $184.62 of $400 (100/650, 300/650): the new run replaced it


def test_mistakes_are_reported_and_nothing_saved(client, org, event):
    overlap = LABOUR + [["wagesonly", "477", "all", "", "", "", "", "Event", "CN", "", "50"]]
    result = say(client, "/splittag --Sep 2026 <<", [rules_file(overlap)])
    assert "is also in rule labour's scope" in notes(result) and "Nothing was saved" in notes(result)
    assert not SplitRun.objects.filter(status="open").exists()

    too_much = [["big", "469-477", "journals", "", "", "", "", "Event", "ABBA", "700", ""]]
    assert "more than the 600.00 in scope" in notes(say(client, "/splittag --Sep 2026 <<", [rules_file(too_much)]))

    bad = [["all", "477", "", "", "", "", "", "Event", "ABBA", "100", "20"],
           ["pct", "477", "bills", "", "", "", "", "Event", "ABBA", "", "80"],
           ["pct", "477", "bills", "", "", "", "", "Event", "ADF", "", "30"],
           ["oops", "477", "bills", "", "", "", "", "Event", "NOPE", "", "10"]]
    message = error(client, "/splittag --Sep 2026 <<", [rules_file(bad)])
    assert '"all" can\'t be a rule name' in message
    assert "more than 100%" in message and '"NOPE" isn\'t one Event value' in message

    assert "Give the period" in error(client, "/splittag <<", [rules_file(LABOUR)])
    assert 'No saved split rule called "nope"' in error(client, "/splittag --rule nope --Sep 2026")


def test_first_time_gets_the_template_and_the_next_command(client, org, event):
    result = say(client, "/splittag --Sep 2026")
    assert result["link_text"] == "Download the template (splittag-template.csv)"
    assert "Filling in the template" in tables(result)
    assert "/splittag --September 2026 << rules.csv" in notes(result)
    body = client.get(result["link"]).content.decode("utf-8-sig")
    assert body.splitlines()[0] == ",".join(HEAD)
    assert say(client, "/splittag --template")["link"] == result["link"]

    say(client, "/splittag --Sep 2026 <<", [rules_file(LABOUR)])
    say(client, "approve")
    result = say(client, "/splittag --Oct 2026")
    assert "Saved rules" in tables(result) and "/splittag --labour --October 2026" in notes(result)


def test_fixed_amounts_are_exact_to_the_cent_on_awkward_totals(client, org, event, source):
    lines = [{"Description": f"Shift {n}", "AccountCode": "477", "LineAmount": amount}
             for n, amount in enumerate(["1234.57", "987.65", "3333.33", "0.01", "54321.09"], 1)]
    lines.append({"Description": "Accrual", "AccountCode": "800", "LineAmount": "-59876.65"})
    source.raw_manual_journals = lambda start, end: [{"ManualJournalID": "mj-odd", "Status": "POSTED",
                                                      "DateString": "2026-09-30T00:00:00", "JournalLines": lines}]
    rows = [["odd", "477", "journals", "", "", "", "", "Event", "ABBA", "100", ""],
            ["odd", "477", "journals", "", "", "", "", "Event", "ADF", "300", ""],
            ["odd", "477", "journals", "", "", "", "", "Event", "CN", "", "25"]]
    result = say(client, "/splittag --Sep 2026 <<", [rules_file(rows)])
    table = tables(result)["odd: 477 · journals · 5 line(s) · September 2026"]
    assert [r[3] for r in table] == ["100.00", "300.00", "14,869.16", "44,607.49", "59,876.65"]
    say(client, "approve")
    source.profit_and_loss = lambda start, end: ProfitAndLoss(by_account={"a-wages": D("59876.65")},
                                                              category_of={"a-wages": "expenses"},
                                                              net_profit=D("-59876.65"))
    pl = next(b for b in say(client, "/pl --event --Sep 2026")["blocks"] if b["type"] == "table")
    wages = next(r["cells"] for r in pl["rows"] if r["cells"][1] == "Wages")
    # the $250 labour-hire bill on Wages isn't a journal, so it stays Unassigned
    assert wages[2:7] == ["100.00", "300.00", "14,869.16", "44,857.49", "60,126.65"]



def test_pl_warns_about_a_split_waiting_for_approval(client, org, event):
    say(client, "/splittag --Sep 2026 <<", [rules_file(LABOUR)])
    result = say(client, "/pl --event --2026-09")
    assert "Your /splittag preview (labour) isn't approved yet" in notes(result)
    assert say(client, "approve")["title"] == "Split saved: 2 line(s) tagged"  # still waiting after the report
    result = say(client, "/pl --event --2026-09")
    assert "isn't approved yet" not in notes(result)
    assert "Unassigned" in [c["label"] for c in next(b for b in result["blocks"] if b["type"] == "table")["columns"]]
    say(client, "/splittag --Sep 2026 <<", [rules_file(LABOUR)])
    assert "isn't approved" not in notes(say(client, "/pl --event --Oct 2026"))  # a different period


def test_pl_by_event_downloads_as_excel(client, org, event):
    from openpyxl import load_workbook

    say(client, "/splittag --Sep 2026 <<", [rules_file(LABOUR)])
    say(client, "approve")
    response = client.post("/console/run/", {"command": "/pl --event --2026-09 >> profit-and-loss-by-event.xlsx"})
    assert response["Content-Disposition"] == 'attachment; filename="profit-and-loss-by-event.xlsx"'
    sheet = load_workbook(io.BytesIO(response.content)).worksheets[-1]
    rows = [[c.value for c in row] for row in sheet.iter_rows()]
    header = next(r for r in rows if r and "Account" in r)
    assert header[2:4] == ["ABBA", "ADF ADF Seminar"] and "Difference" in header
    wages = next(r for r in rows if r and r[1] == "Wages")
    assert wages[2] == pytest.approx(66.67)  # a number Excel can add up, not text
