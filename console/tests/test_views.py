import io
from unittest import mock

import pytest
from django.contrib.auth import get_user_model
from openpyxl import load_workbook

from core.models import AuditLog, Organisation

from .fake_xero import fake_source


@pytest.fixture
def user(db):
    return get_user_model().objects.create_user("finance", password="pw")


@pytest.fixture
def orgs(user):
    demo = Organisation.objects.create(name="Demo Events")
    other = Organisation.objects.create(name="Other Venue")
    demo.members.add(user)
    other.members.add(user)
    return demo, other


@pytest.fixture
def client(client, user):
    client.force_login(user)
    return client


@pytest.fixture(autouse=True)
def fake(monkeypatch):
    source = fake_source()
    monkeypatch.setattr("console.views.source_for", lambda organisation: source)
    return source


def post(client, command, **extra):
    return client.post("/console/run/", {"command": command, **extra})


def test_console_requires_login(db, client):
    client.logout()
    assert client.get("/console/").status_code == 302
    assert client.post("/console/run/", {"command": "/help"}).status_code == 302


def test_root_redirects_to_console(client):
    assert client.get("/").headers["Location"] == "/console/"


def test_console_page_renders(client, orgs):
    page = client.get("/console/")
    assert page.status_code == 200
    assert b"Demo Events" in page.content


def test_run_returns_json_and_audits(client, orgs):
    data = post(client, "/status --September 2026 --revenue").json()
    assert data["ok"] and data["result"]["title"] == "Performance against budget and last year"
    row = data["result"]["blocks"][0]["rows"][0]
    assert row["cells"][2] == "10,000.00" and row["cells"][4] == "(2,000.00)" and row["negative"][4]
    log = AuditLog.objects.get()
    assert log.action == AuditLog.Action.COMMAND and log.succeeded and log.organisation == orgs[0]
    assert log.input == "/status --September 2026 --revenue"


def test_user_errors_are_shown_and_audited(client, orgs):
    data = post(client, "/nonsense").json()
    assert not data["ok"] and "no /nonsense command" in data["error"]
    data = post(client, "/status << budget.csv").json()
    assert "doesn't take an input file" in data["error"]
    assert AuditLog.objects.filter(succeeded=False).count() == 2


def test_unexpected_errors_are_hidden(client, orgs, fake):
    with mock.patch.object(type(fake), "bank_balances", side_effect=RuntimeError("secret detail")):
        resp = post(client, "/cash")
    assert resp.status_code == 500
    assert "secret" not in resp.json()["error"]


def test_export_xlsx(client, orgs):
    resp = post(client, "/pl --September 2026 >> ../../pl sep.xlsx")
    assert resp["Content-Disposition"] == 'attachment; filename="pl sep.xlsx"'
    sheet = load_workbook(io.BytesIO(resp.content)).active
    values = [[c.value for c in row] for row in sheet.iter_rows()]
    assert values[0][0] == "Profit and loss: September 2026"
    sales = next(r for r in values if r[1] == "Sales")
    assert sales[0] == "200" and sales[2] == 10000.0
    assert sheet.cell(row=values.index(sales) + 1, column=1).number_format == "@"


def test_export_csv_and_bad_extension(client, orgs):
    resp = post(client, "/bills >> bills.csv")
    text = resp.content.decode("utf-8-sig")
    assert text.splitlines()[0] == "Due,Supplier,Number,Status,Currency,Amount due"
    assert "2026-10-05,Acme Foods,B-1,Overdue,AUD,300.00" in text
    assert "Save as .xlsx or .csv" in post(client, "/bills >> bills.pdf").json()["error"]


def test_orgs_and_use_switch_organisation(client, orgs):
    data = post(client, "/use other").json()
    assert data["result"]["title"] == "Now using Other Venue" and data["organisation"] == "Other Venue"
    rows = post(client, "/orgs").json()["result"]["blocks"][0]["rows"]
    assert [r["cells"][:2] for r in rows] == [["", "Demo Events"], ["*", "Other Venue"]]
    assert "matches" in post(client, "/use e").json()["error"]


def test_other_users_organisation_is_not_selectable(client, orgs):
    stranger = Organisation.objects.create(name="Stranger Pty Ltd")
    session = client.session
    session["organisation_id"] = stranger.pk
    session.save()
    assert post(client, "/help").json()["organisation"] == "Demo Events"


def test_no_organisation_yet(client):
    data = post(client, "/status").json()
    assert not data["ok"] and "/orgs" in data["error"]
    assert post(client, "/help").json()["ok"]
